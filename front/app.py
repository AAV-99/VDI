"""
app.py — Interfaz Streamlit para el framework multi-agente VDI 2206.

Flujo:
  1. El usuario escribe el contexto del problema y (opcional) adjunta documentos.
  2. Se genera input/project_context.txt (contexto + texto extraído de los adjuntos).
  3. Se ejecutan los agentes (simulados o reales).
  4. Se muestran los .md de output/ en pantalla.

Ejecutar (desde la raíz del repo, con el entorno virtual activado):
  streamlit run front/app.py
"""
from __future__ import annotations

import io
import os
import re
import shutil
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path

import streamlit as st

from mermaid_render import render_markdown_with_mermaid
from mock_flow import DOSSIER_NAME, run_mock_flow

ROOT = Path(__file__).resolve().parent
INPUT_DIR = ROOT / "input"
INPUT_FILE = INPUT_DIR / "project_context.txt"
MOCK_OUTPUT_DIR = ROOT / "output"

# --- Flujo real de CrewAI: vive en la raíz del repo (carpeta padre de front/) ---
REPO_ROOT = ROOT.parent
REAL_OUTPUT_DIR = REPO_ROOT / "output"  # ahí escriben los agentes (vdicrew.py)
# Se usa el mismo Python que corre Streamlit: funciona igual con pip/venv o con uv
REAL_FLOW_CMD = [sys.executable, "-m", "vdi.main"]
CONSOLE_LOG = REAL_OUTPUT_DIR / "_consola.log"  # salida completa de la corrida real
PID_FILE = REAL_OUTPUT_DIR / "_corrida.pid"  # PID de la corrida real en curso (se borra al terminar)
ENV_FILE = REPO_ROOT / ".env"  # aquí va GEMINI_API_KEY
CONTEXT_ENV_VAR = "VDI_CONTEXT_FILE"  # src/vdi/main.py lee el .txt desde esta variable
REPORT_PATTERN = re.compile(r"^\d\d_.*\.md$")  # 00_..., 01_..., etc.

ALLOWED_TYPES = ["txt", "md", "csv", "pdf", "docx"]


# ----------------------------------------------------------------------------
# Extracción de texto de documentos adjuntos
# ----------------------------------------------------------------------------
def extract_text(uploaded_file) -> str:
    """Devuelve el texto de un archivo subido (txt, md, csv, pdf, docx)."""
    name = uploaded_file.name
    ext = Path(name).suffix.lower()
    data = uploaded_file.getvalue()
    print(f"[UI] Extrayendo texto de {name} ({len(data)} bytes, tipo {ext})")

    if ext in (".txt", ".md", ".csv"):
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            print(f"[UI]    {name}: no es UTF-8, se usa latin-1")
            text = data.decode("latin-1")

    elif ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
        print(f"[UI]    {name}: {len(pages)} páginas")
        text = "\n\n".join(pages)

    elif ext == ".docx":
        import docx

        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(parts)

    else:
        raise ValueError(f"Tipo de archivo no soportado: {ext}")

    print(f"[UI]    {name}: {len(text)} caracteres extraídos")
    return text.strip()


# ----------------------------------------------------------------------------
# Construcción del .txt de entrada para los agentes
# ----------------------------------------------------------------------------
def build_input_txt(system_type, scope, requirements, cost, extra, documents) -> str:
    """documents: lista de tuplas (nombre, texto)."""
    lines = [
        "=== CONTEXTO DEL PROYECTO VDI 2206 ===",
        f"Generado: {datetime.now():%Y-%m-%d %H:%M:%S}",
        "",
        "## TIPO DE SISTEMA",
        system_type.strip(),
        "",
        "## ALCANCE DEL PROYECTO",
        scope.strip(),
        "",
        "## REQUERIMIENTOS PRIMARIOS",
        requirements.strip() or "(no especificados)",
        "",
        "## COSTO OBJETIVO",
        cost.strip() or "(no especificado)",
        "",
        "## CONTEXTO ADICIONAL",
        extra.strip() or "(ninguno)",
        "",
    ]
    if documents:
        lines += ["=== DOCUMENTOS ADJUNTOS ===", ""]
        for i, (name, text) in enumerate(documents, start=1):
            lines += [
                f"--- DOCUMENTO {i}: {name} ---",
                text or "(sin texto extraíble)",
                f"--- FIN DOCUMENTO {i} ---",
                "",
            ]
    return "\n".join(lines)


def report_files(output_dir: Path) -> list[Path]:
    """Reportes de los agentes (00_..., 01_...) presentes en output_dir."""
    return sorted(f for f in output_dir.glob("*.md") if REPORT_PATTERN.match(f.name))


def _move_with_retry(src: Path, dst: Path, attempts: int = 5, wait: float = 0.6) -> str:
    """Mueve src a dst. Si Windows/OneDrive lo tiene bloqueado, reintenta y al final lo copia.

    Devuelve "movido", "copiado" o "omitido". Nunca lanza excepción.
    """
    for intento in range(1, attempts + 1):
        try:
            src.replace(dst)
            return "movido"
        except PermissionError as exc:
            print(f"[UI]    {src.name} bloqueado (intento {intento}/{attempts}): {exc.strerror}")
            time.sleep(wait)
        except OSError as exc:
            print(f"[UI]    {src.name}: error al mover ({exc}); se intenta copiar")
            break
    try:
        shutil.copy2(src, dst)  # copiar sí se permite aunque el archivo esté abierto
        return "copiado"
    except OSError as exc:
        print(f"[UI]    {src.name}: tampoco se pudo copiar ({exc}); se deja donde está")
        return "omitido"


def archive_previous_outputs(output_dir: Path) -> None:
    """Mueve los resultados de la corrida anterior a output/_anteriores/<fecha>/ (no los borra)."""
    old_files = [f for f in output_dir.iterdir() if f.is_file() and f.name != PID_FILE.name]
    if not old_files:
        print("[UI] No hay resultados anteriores que archivar")
        return
    dest = output_dir / "_anteriores" / f"{datetime.now():%Y%m%d_%H%M%S}"
    dest.mkdir(parents=True, exist_ok=True)
    resumen = {"movido": 0, "copiado": 0, "omitido": 0}
    for old in old_files:
        resultado = _move_with_retry(old, dest / old.name)
        resumen[resultado] += 1
        if resultado != "movido":
            print(f"[UI]    {old.name}: {resultado}")
    print(f"[UI] Archivado en {dest}: {resumen}")
    if resumen["copiado"] or resumen["omitido"]:
        # Los que quedaron bloqueados se borran si se puede, para que no se mezclen con la corrida nueva
        for old in old_files:
            if old.exists():
                try:
                    old.unlink()
                except OSError:
                    print(f"[UI]    {old.name} sigue bloqueado; se sobrescribirá en la corrida nueva")


def _pid_alive(pid: int) -> bool:
    """True si el proceso sigue vivo. En Windows NO se usa os.kill(pid, 0): eso lo mataría."""
    if pid <= 0:
        return False
    if os.name == "nt":
        import ctypes

        PROCESS_QUERY_LIMITED_INFORMATION = 0x1000
        STILL_ACTIVE = 259
        kernel32 = ctypes.windll.kernel32
        handle = kernel32.OpenProcess(PROCESS_QUERY_LIMITED_INFORMATION, False, pid)
        if not handle:
            return False
        try:
            code = ctypes.c_ulong()
            ok = kernel32.GetExitCodeProcess(handle, ctypes.byref(code))
            return bool(ok) and code.value == STILL_ACTIVE
        finally:
            kernel32.CloseHandle(handle)
    try:
        os.kill(pid, 0)
        return True
    except OSError:
        return False


def running_flow_pid() -> int | None:
    """PID de una corrida real que siga en curso (p. ej. tras recargar la página), o None."""
    if not PID_FILE.exists():
        return None
    try:
        pid = int(PID_FILE.read_text(encoding="utf-8").strip())
    except (ValueError, OSError):
        pid = 0
    if _pid_alive(pid):
        print(f"[UI] Hay una corrida real en curso con PID {pid}")
        return pid
    print(f"[UI] PID {pid} ya no está activo: se limpia {PID_FILE.name}")
    try:
        PID_FILE.unlink()
    except OSError:
        pass
    return None


def stop_flow(pid: int) -> None:
    """Detiene la corrida (y sus subprocesos, p. ej. el servidor MCP)."""
    print(f"[UI] Deteniendo corrida con PID {pid}")
    if os.name == "nt":
        r = subprocess.run(["taskkill", "/PID", str(pid), "/T", "/F"], capture_output=True, text=True)
        print(f"[UI]    taskkill: {(r.stdout or r.stderr).strip()}")
    else:
        import signal

        try:
            os.kill(pid, signal.SIGTERM)
        except OSError as exc:
            print(f"[UI]    no se pudo detener: {exc}")
    time.sleep(1.5)  # Windows tarda un momento en liberar los archivos
    try:
        PID_FILE.unlink()
    except OSError:
        pass


def run_real_flow(on_log=None) -> str:
    """Lanza el flujo real de CrewAI desde la raíz del repo y devuelve su salida de consola.

    on_log(texto): se llama con el log acumulado cada vez que llega una línea nueva.
    """
    env = os.environ.copy()
    env[CONTEXT_ENV_VAR] = str(INPUT_FILE)
    # CrewAI imprime emojis/cajas: sin esto, en Windows la salida redirigida falla con cp1252
    env["PYTHONIOENCODING"] = "utf-8"
    env["PYTHONUTF8"] = "1"
    env["PYTHONUNBUFFERED"] = "1"
    # El paquete vdi vive en src/: así no hace falta instalar el proyecto
    env["PYTHONPATH"] = os.pathsep.join(filter(None, [str(REPO_ROOT / "src"), env.get("PYTHONPATH", "")]))
    # Variables del .env de la raíz (no pisan las que ya estén definidas en la terminal)
    if ENV_FILE.exists():
        from dotenv import dotenv_values

        for key, value in dotenv_values(ENV_FILE).items():
            if value is not None:
                env.setdefault(key, value)
    else:
        print(f"[UI] AVISO: no existe {ENV_FILE}")
    print(f"[UI]    GEMINI_API_KEY definida: {bool(env.get('GEMINI_API_KEY') or env.get('GOOGLE_API_KEY'))}")
    print(f"[UI] Ejecutando flujo real: {' '.join(REAL_FLOW_CMD)}")
    print(f"[UI]    carpeta de trabajo: {REPO_ROOT}")
    print(f"[UI]    {CONTEXT_ENV_VAR}={INPUT_FILE}")

    # La salida va a un ARCHIVO y no a un PIPE. Con PIPE, si Streamlit recarga la página a mitad
    # de la corrida, nadie vuelve a leer el PIPE, su buffer se llena y el flujo se congela sin error
    # (así se quedaban sin escribir los últimos reportes). Con el archivo, los agentes siguen
    # aunque la vista se interrumpa, y la consola completa queda en output/_consola.log.
    CONSOLE_LOG.parent.mkdir(parents=True, exist_ok=True)
    console_log = CONSOLE_LOG
    try:
        salida = open(console_log, "w", encoding="utf-8")
    except PermissionError:
        # _consola.log bloqueado (OneDrive u otro proceso): se usa un nombre con fecha
        console_log = CONSOLE_LOG.with_name(f"_consola_{datetime.now():%Y%m%d_%H%M%S}.log")
        print(f"[UI]    {CONSOLE_LOG.name} bloqueado; se usa {console_log.name}")
        salida = open(console_log, "w", encoding="utf-8")
    with salida:
        process = subprocess.Popen(
            REAL_FLOW_CMD,
            cwd=REPO_ROOT,
            env=env,
            stdout=salida,
            stderr=subprocess.STDOUT,
        )
    PID_FILE.write_text(str(process.pid), encoding="utf-8")
    print(f"[UI]    PID de la corrida: {process.pid} (guardado en {PID_FILE.name})")
    print(f"[UI]    consola completa en {console_log}")
    lines = []
    with open(console_log, encoding="utf-8", errors="replace") as lector:
        while True:
            fragmento = lector.read()
            if fragmento:
                print(fragmento, end="", flush=True)  # también queda en la terminal donde corre Streamlit
                lines.append(fragmento)
                if on_log:
                    on_log("".join(lines))
            elif process.poll() is not None:
                break
            else:
                time.sleep(0.5)
    returncode = process.wait()
    try:
        PID_FILE.unlink()
    except OSError:
        pass
    print(f"[UI] Flujo real terminó con código {returncode}")
    log = "".join(lines)
    if returncode != 0:
        raise RuntimeError(f"El flujo terminó con código {returncode}:\n{log[-3000:]}")
    return log


# ----------------------------------------------------------------------------
# Interfaz
# ----------------------------------------------------------------------------
st.set_page_config(page_title="VDI 2206 — Revisión de Diseño", page_icon="⚙️", layout="wide")
st.title("⚙️ VDI 2206 — Revisión de Diseño Mecatrónico")
st.caption("Describe el problema, adjunta documentos de soporte y ejecuta la tripulación de agentes.")

with st.sidebar:
    st.header("Configuración")
    mode = st.radio(
        "Modo de ejecución",
        ["Simulación (mock)", "Agentes reales (CrewAI)"],
        help="La simulación no llama al modelo: sirve para probar la interfaz.",
    )
    OUTPUT_DIR = MOCK_OUTPUT_DIR if mode.startswith("Simulación") else REAL_OUTPUT_DIR
    if not mode.startswith("Simulación"):
        # Verificación previa: aquí no se instala nada, solo se avisa qué falta
        import importlib.util

        missing = [pkg for pkg in ("crewai", "google.genai", "dotenv") if importlib.util.find_spec(pkg) is None]
        print(f"[UI] Python en uso: {sys.executable}")
        print(f"[UI] Paquetes faltantes para el flujo real: {missing or 'ninguno'}")
        if missing:
            st.error(
                "Faltan paquetes en este entorno: " + ", ".join(missing) + ". Instálalos antes de correr:\n\n"
                '`pip install "crewai[google-genai,tools]>=1.15.24,<2.0.0"`'
            )
        if not ENV_FILE.exists():
            st.error(f"No existe `{ENV_FILE.name}` en la raíz del proyecto con GEMINI_API_KEY.")
    st.divider()
    st.markdown(f"**Input:** `{INPUT_FILE.relative_to(REPO_ROOT)}`")
    st.markdown(f"**Output:** `{OUTPUT_DIR.relative_to(REPO_ROOT)}/`")

st.subheader("1. Contexto del problema")
system_type = st.text_input(
    "Tipo de sistema *",
    placeholder="Ej.: Efector final (gripper) para manipulación de láminas de acero",
)
scope = st.text_area(
    "Alcance del proyecto *",
    height=160,
    placeholder="Qué se va a diseñar, dónde opera, con qué equipos se integra, qué piezas maneja...",
)
requirements = st.text_area(
    "Requerimientos primarios",
    height=120,
    placeholder="Ej.: Masa <= 1.5 kg. Tiempo de ciclo <= 4.0 s. Alimentación 24V DC...",
)
cost = st.text_input(
    "Costo objetivo",
    placeholder="Ej.: Costo objetivo de fabricación <= 1000 USD",
)
extra = st.text_area(
    "Contexto adicional",
    height=100,
    placeholder="Restricciones, normas aplicables, presupuesto, antecedentes...",
)

st.subheader("2. Documentos de soporte (opcional)")
uploads = st.file_uploader(
    "Adjunta especificaciones, fichas técnicas, normas, etc.",
    type=ALLOWED_TYPES,
    accept_multiple_files=True,
)

st.subheader("3. Ejecutar")
busy_pid = None if mode.startswith("Simulación") else running_flow_pid()
if busy_pid:
    st.warning(
        f"Hay una corrida de agentes todavía en curso (PID {busy_pid}), probablemente porque la página se "
        "recargó o el computador se bloqueó mientras corría. Sus resultados siguen llegando a `output/`. "
        "Espera a que termine (recarga la página en unos minutos) o detenla para lanzar una nueva."
    )
    if st.button("⏹️ Detener la corrida en curso"):
        stop_flow(busy_pid)
        st.rerun()
if st.button("Generar input y ejecutar agentes", type="primary", disabled=bool(busy_pid)):
    if not system_type.strip() or not scope.strip():
        st.error("Completa al menos el tipo de sistema y el alcance del proyecto.")
        st.stop()

    # --- 3.1 Extraer texto de los adjuntos ---
    documents = []
    for up in uploads or []:
        try:
            documents.append((up.name, extract_text(up)))
        except Exception as exc:  # un adjunto dañado no debe tumbar todo
            print(f"[UI] ERROR extrayendo {up.name}: {exc}")
            st.warning(f"No se pudo leer **{up.name}**: {exc}")
    for name, text in documents:
        if not text:
            st.warning(f"**{name}** no tiene texto extraíble (¿PDF escaneado?).")

    # --- 3.2 Escribir el .txt ---
    input_txt = build_input_txt(system_type, scope, requirements, cost, extra, documents)
    INPUT_DIR.mkdir(parents=True, exist_ok=True)
    INPUT_FILE.write_text(input_txt, encoding="utf-8")
    print(f"[UI] Input escrito en {INPUT_FILE} ({len(input_txt)} caracteres, {len(documents)} adjuntos)")
    st.session_state["input_txt"] = input_txt

    # --- 3.3 Apartar resultados de la corrida anterior ---
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    st.session_state["output_dir"] = str(OUTPUT_DIR)
    st.session_state.pop("run_log", None)
    if mode.startswith("Simulación"):
        for old in OUTPUT_DIR.glob("*.md"):
            old.unlink()
            print(f"[UI] Eliminado resultado anterior: {old.name}")
    else:
        archive_previous_outputs(OUTPUT_DIR)  # una corrida real tarda: no se borra, se archiva

    # --- 3.4 Ejecutar agentes ---
    try:
        if mode.startswith("Simulación"):
            progress = st.progress(0.0, text="Iniciando agentes...")
            run_mock_flow(
                INPUT_FILE,
                OUTPUT_DIR,
                on_step=lambda i, total, label: progress.progress(i / total, text=f"{label} ({i}/{total})"),
            )
            progress.empty()
        else:
            st.info("No cambies nada ni recargues la página mientras corre: eso interrumpe la vista. Los agentes siguen corriendo; los resultados quedan en output/ y la consola completa en output/_consola.log.")
            with st.spinner("Ejecutando agentes reales... esto puede tardar varios minutos."):
                live_log = st.empty()
                st.session_state["run_log"] = run_real_flow(
                    on_log=lambda log: live_log.code(log[-3000:], language="text")
                )
                live_log.empty()
        st.session_state["run_ok"] = True
        st.success("Ejecución completada.")
    except Exception as exc:
        print(f"[UI] ERROR en la ejecución: {exc}")
        st.session_state["run_ok"] = False
        st.error(f"Falló la ejecución de los agentes:\n\n```\n{exc}\n```")

# ----------------------------------------------------------------------------
# Resultados (persisten entre recargas gracias a session_state)
# ----------------------------------------------------------------------------
if "input_txt" in st.session_state:
    with st.expander("📄 Ver el .txt enviado a los agentes"):
        st.download_button(
            "Descargar project_context.txt",
            st.session_state["input_txt"],
            file_name="project_context.txt",
            mime="text/plain",
        )
        st.code(st.session_state["input_txt"], language="text")

if st.session_state.get("run_log"):
    with st.expander("🖥️ Log de consola de los agentes"):
        st.code(st.session_state["run_log"][-10000:], language="text")

# Los resultados se leen directamente de la carpeta del modo elegido. Así se ven aunque
# la página se haya recargado o la vista se haya interrumpido durante la corrida.
results_dir = OUTPUT_DIR
md_files = report_files(results_dir) if results_dir.exists() else []
print(f"[UI] Reportes encontrados en {results_dir}: {[f.name for f in md_files]}")

if md_files:
    st.subheader("4. Resultados")
    last_change = datetime.fromtimestamp(max(f.stat().st_mtime for f in md_files))
    st.caption(f"Carpeta `{results_dir.relative_to(REPO_ROOT)}/` · última actualización: {last_change:%Y-%m-%d %H:%M}")
    # El dossier (00_...) queda primero por orden alfabético
    tabs = st.tabs(["📘 Dossier" if f.name == DOSSIER_NAME else f.stem for f in md_files])
    for tab, md_file in zip(tabs, md_files):
        with tab:
            content = md_file.read_text(encoding="utf-8")
            st.download_button(
                f"Descargar {md_file.name}",
                content,
                file_name=md_file.name,
                mime="text/markdown",
                key=f"dl_{md_file.name}",
            )
            # Los diagramas Mermaid se sanean, se validan y, si aun así fallan, se muestran como texto
            render_markdown_with_mermaid(content)
